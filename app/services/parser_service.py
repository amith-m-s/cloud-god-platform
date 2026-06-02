from __future__ import annotations

import io
import structlog

logger = structlog.get_logger("parser_service")

def extract_text_from_blob(blob: bytes, filename: str) -> str:
    """
    Extracts plain text from a document binary payload based on its file extension.
    Supports PDF, DOCX, TXT, MD, and falls back to plain UTF-8 decoding.
    """
    if not blob:
        return ""

    ext = filename.split(".")[-1].lower() if "." in filename else ""
    logger.info("parsing_started", filename=filename, ext=ext, size=len(blob))

    try:
        if ext == "pdf":
            return _parse_pdf(blob)
        elif ext == "docx":
            return _parse_docx(blob)
        elif ext in ("txt", "md", "csv", "json", "xml"):
            return _parse_plain_text(blob)
        else:
            # General fallback
            return _parse_plain_text(blob)
    except Exception as exc:
        logger.error("parsing_failed", filename=filename, error=str(exc))
        # Final safety fallback: try standard utf-8 decoding
        return blob.decode("utf-8", errors="ignore")


def _parse_pdf(blob: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(blob))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(text)
        extracted = "\n".join(pages_text)
        logger.info("pdf_parsed_successfully", pages=len(reader.pages), chars=len(extracted))
        return extracted
    except ImportError:
        logger.error("pypdf_not_installed_falling_back")
        return blob.decode("utf-8", errors="ignore")


def _parse_docx(blob: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(blob))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        # Also extract from tables for complete coverage
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        paragraphs.append(cell.text)
        extracted = "\n".join(paragraphs)
        logger.info("docx_parsed_successfully", paragraphs=len(paragraphs), chars=len(extracted))
        return extracted
    except ImportError:
        logger.error("docx_not_installed_falling_back")
        return blob.decode("utf-8", errors="ignore")


def _parse_plain_text(blob: bytes) -> str:
    # Try different encodings
    encodings = ["utf-8", "latin-1", "cp1252", "utf-16"]
    for enc in encodings:
        try:
            decoded = blob.decode(enc)
            logger.info("text_parsed_successfully", encoding=enc, chars=len(decoded))
            return decoded
        except UnicodeDecodeError:
            continue
    return blob.decode("utf-8", errors="ignore")
